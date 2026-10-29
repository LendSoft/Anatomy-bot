"""Извлечение текста и иллюстраций из учебных материалов.

Разбор PDF идёт через PyMuPDF: он отдаёт и текст, и координаты строк, и
границы картинок за один проход. Координаты обязательны — без них
невозможно понять, какая подпись относится к какому рисунку и что
написано внутри изображения (см. app/rag/figures.py).

Порядок для PDF:
  1. текстовый слой — быстро и достаточно для книг, свёрстанных в вёрстке;
  2. если текста почти нет, документ считается сканом и идёт в OCR;
  3. иллюстрации ищутся в обоих случаях, независимо от текстового слоя.

Номера страниц сохраняются на каждом шаге: без них не получится указать
источник так, как требует ТЗ, и не получится сослаться на страницу с
рисунком.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

from app.core.logging import get_logger
from app.rag import figures as figures_mod
from app.rag.figures import ExtractedFigure, FigureKind

log = get_logger(__name__)

# Ниже этого числа символов на страницу считаем, что текстового слоя нет
MIN_CHARS_PER_PAGE = 80

# Разрешение растеризации для OCR. 300 dpi — компромисс: ниже падает
# качество распознавания русских диакритик и мелких выносок, выше
# страница сканируется ощутимо дольше без выигрыша в точности
OCR_DPI = 300


@dataclass
class PageText:
    page_number: int
    text: str
    # Заголовки, найденные на странице — используются как метка фрагмента
    headings: list[str] = field(default_factory=list)
    # Иллюстрации этой страницы
    figures: list[ExtractedFigure] = field(default_factory=list)


@dataclass
class ParsedDocument:
    pages: list[PageText]
    used_ocr: bool = False
    parser: str = "pymupdf"

    @property
    def total_chars(self) -> int:
        return sum(len(p.text) for p in self.pages)

    @property
    def figures(self) -> list[ExtractedFigure]:
        return [figure for page in self.pages for figure in page.figures]


class ParsingError(RuntimeError):
    pass


class ParsingCancelled(RuntimeError):
    pass


def _check_cancelled(should_cancel: Callable[[], bool] | None) -> None:
    if should_cancel is not None and should_cancel():
        raise ParsingCancelled("Индексация отменена")


# ------------------------------------------------------------------ утилиты

_HYPHEN_BREAK = re.compile(r"(\w)-\s*\n\s*(\w)")
_MULTI_SPACE = re.compile(r"[ \t]+")
_MULTI_NEWLINE = re.compile(r"\n{3,}")
# Строка целиком из заглавных или похожая на нумерованный заголовок раздела
_HEADING = re.compile(
    r"^(?:[А-ЯЁA-Z][А-ЯЁA-Z\s\-,\.]{6,80}|"
    r"(?:\d+\.){1,3}\s+[А-ЯЁA-Z][^\n]{3,80})$"
)


def clean_text(raw: str) -> str:
    """Убирает переносы по слогам и лишние пробелы.

    В учебниках слова регулярно разорваны переносом в конце строки —
    без склейки «плече-вая» не найдётся ни поиском, ни эмбеддингом.
    """
    text = raw.replace("\xa0", " ").replace("­", "")
    text = _HYPHEN_BREAK.sub(r"\1\2", text)
    text = _MULTI_SPACE.sub(" ", text)
    text = _MULTI_NEWLINE.sub("\n\n", text)
    return text.strip()


def detect_headings(text: str) -> list[str]:
    return [
        line.strip()
        for line in text.split("\n")
        if _HEADING.match(line.strip())
    ]


# -------------------------------------------------------------------- PDF


def open_pdf(path: Path):
    """Открывает PDF. Вынесено отдельно: тем же документом пользуется
    вырезание предпросмотра иллюстрации."""
    import pymupdf

    try:
        return pymupdf.open(str(path))
    except Exception as exc:  # noqa: BLE001
        raise ParsingError(f"Не удалось открыть PDF: {exc}") from exc


def parse_pdf_text_layer(
    doc, should_cancel: Callable[[], bool] | None = None
) -> ParsedDocument:
    pages: list[PageText] = []

    for i in range(len(doc)):
        _check_cancelled(should_cancel)
        page = doc[i]
        page_number = i + 1
        try:
            raw = page.get_text("text") or ""
        except Exception as exc:  # noqa: BLE001 — одна битая страница не должна ронять книгу
            log.warning("pdf_page_extract_failed", page=page_number, error=str(exc))
            raw = ""

        text = clean_text(raw)
        try:
            page_figures = figures_mod.extract_page_figures(page, page_number)
        except Exception as exc:  # noqa: BLE001 — рисунки необязательны
            log.warning("pdf_page_figures_failed", page=page_number, error=str(exc))
            page_figures = []

        pages.append(
            PageText(
                page_number=page_number,
                text=text,
                headings=detect_headings(text),
                figures=page_figures,
            )
        )

    return ParsedDocument(pages=pages, parser="pymupdf")


def _ocr_page(page, languages: str) -> str:
    """Распознаёт одну страницу.

    Страница растеризуется средствами PyMuPDF, минуя pdf2image и
    poppler: лишняя зависимость в контейнере не нужна, когда в проекте
    уже есть библиотека, умеющая отдавать пиксели.
    """
    import io

    import pytesseract
    from PIL import Image

    pixmap = page.get_pixmap(dpi=OCR_DPI)
    with Image.open(io.BytesIO(pixmap.tobytes("png"))) as image:
        return pytesseract.image_to_string(image, lang=languages)


def parse_pdf_ocr(
    doc,
    languages: str = "rus+eng",
    should_cancel: Callable[[], bool] | None = None,
) -> ParsedDocument:
    """Распознавание отсканированных страниц.

    Медленно (единицы секунд на страницу), поэтому вызывается только
    когда текстового слоя действительно нет.
    """
    pages: list[PageText] = []
    total = len(doc)

    for i in range(total):
        _check_cancelled(should_cancel)
        page = doc[i]
        page_number = i + 1
        try:
            raw = _ocr_page(page, languages)
        except Exception as exc:  # noqa: BLE001
            log.warning("ocr_page_failed", page=page_number, error=str(exc))
            raw = ""

        text = clean_text(raw)
        try:
            page_figures = figures_mod.extract_page_figures(page, page_number)
        except Exception as exc:  # noqa: BLE001 — рисунки необязательны
            log.warning("ocr_page_figures_failed", page=page_number, error=str(exc))
            page_figures = []

        # У скана весь лист — одна картинка. Текст распознан отдельно, так
        # что подписи структур берутся из него, а не из границ изображения.
        for figure in page_figures:
            if figure.kind == FigureKind.page and text:
                figure.title = figure.title or (detect_headings(text) or [None])[0]

        pages.append(
            PageText(
                page_number=page_number,
                text=text,
                headings=detect_headings(text),
                figures=page_figures,
            )
        )

        if page_number % 25 == 0:
            log.info("ocr_progress", page=page_number, total=total)

    return ParsedDocument(pages=pages, used_ocr=True, parser="pymupdf+tesseract")


def parse_pdf(
    path: Path,
    ocr_enabled: bool = True,
    should_cancel: Callable[[], bool] | None = None,
) -> ParsedDocument:
    doc = open_pdf(path)
    try:
        parsed = parse_pdf_text_layer(doc, should_cancel=should_cancel)

        if not parsed.pages:
            raise ParsingError("PDF не содержит страниц")

        avg_chars = parsed.total_chars / len(parsed.pages)
        if avg_chars >= MIN_CHARS_PER_PAGE:
            return parsed

        if not ocr_enabled:
            raise ParsingError(
                "В PDF нет текстового слоя (похоже на скан), а OCR отключён в настройках"
            )

        log.info(
            "pdf_looks_scanned",
            avg_chars=round(avg_chars, 1),
            pages=len(parsed.pages),
        )
        return parse_pdf_ocr(doc, should_cancel=should_cancel)
    finally:
        doc.close()


# -------------------------------------------------------------------- DOCX

# Подпись к рисунку в методичке — обычным абзацем под картинкой
_DOCX_CAPTION = figures_mod.CAPTION_RE


def parse_docx(
    path: Path, should_cancel: Callable[[], bool] | None = None
) -> ParsedDocument:
    """DOCX не разбит на страницы — разрывы страниц там условны.

    Собираем весь текст, разметку заголовков берём из стилей Word
    (Heading 1..9) и складываем в одну «страницу»: для цитирования
    важнее раздел, чем номер страницы, которого в файле попросту нет.

    Рисунки в DOCX ищутся по подписям и по наличию вложенной картинки в
    абзаце. Номера страницы у них нет, поэтому ссылка на иллюстрацию
    указывает раздел — это всё, на что можно опереться в таком файле.
    """
    import docx

    document = docx.Document(str(path))
    parts: list[str] = []
    headings: list[str] = []
    found: list[ExtractedFigure] = []
    current_heading: str | None = None

    for para in document.paragraphs:
        _check_cancelled(should_cancel)
        text = para.text.strip()
        style = (para.style.name or "").lower()
        # Вложенная картинка видна в XML абзаца как элемент <w:drawing>
        has_image = "graphicData" in para._p.xml or "<w:drawing" in para._p.xml

        if has_image:
            found.append(
                ExtractedFigure(
                    page_number=0,  # страниц в DOCX нет
                    kind=FigureKind.figure,
                    title=current_heading,
                    confidence=0.6,
                )
            )

        if not text:
            continue

        caption_match = _DOCX_CAPTION.match(text)
        if caption_match:
            label = caption_match.group("label").strip()
            rest = caption_match.group("rest").strip()
            # Подпись идёт сразу за своей картинкой — дополняем последнюю
            if found and found[-1].label is None:
                found[-1].label = label
                found[-1].title = rest or current_heading
                found[-1].confidence = 0.9
            else:
                found.append(
                    ExtractedFigure(
                        page_number=0,
                        kind=FigureKind.figure,
                        label=label,
                        title=rest or current_heading,
                        confidence=0.8,
                    )
                )
            parts.append(text)
            continue

        if style.startswith("heading") or style.startswith("заголовок"):
            current_heading = text
            headings.append(text)
            parts.append(f"\n\n## {text}\n")
        else:
            parts.append(text)

    # Таблицы в методичках часто содержат классификации — теряем их зря
    for table in document.tables:
        _check_cancelled(should_cancel)
        rows = [
            " | ".join(cell.text.strip() for cell in row.cells)
            for row in table.rows
            if any(cell.text.strip() for cell in row.cells)
        ]
        if rows:
            parts.append("\n" + "\n".join(rows) + "\n")

    text = clean_text("\n".join(parts))
    return ParsedDocument(
        pages=[
            PageText(
                page_number=1, text=text, headings=headings, figures=found
            )
        ],
        parser="python-docx",
    )


def parse_txt(path: Path) -> ParsedDocument:
    text = clean_text(path.read_text(encoding="utf-8", errors="replace"))
    return ParsedDocument(
        pages=[PageText(page_number=1, text=text, headings=detect_headings(text))],
        parser="plain",
    )


def parse(
    path: Path,
    ocr_enabled: bool = True,
    should_cancel: Callable[[], bool] | None = None,
) -> ParsedDocument:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return parse_pdf(
            path, ocr_enabled=ocr_enabled, should_cancel=should_cancel
        )
    if suffix in (".docx", ".doc"):
        if suffix == ".doc":
            raise ParsingError(
                "Формат .doc не поддерживается — пересохраните файл как .docx"
            )
        return parse_docx(path, should_cancel=should_cancel)
    if suffix in (".txt", ".md"):
        return parse_txt(path)
    raise ParsingError(f"Неподдерживаемый формат файла: {suffix}")
